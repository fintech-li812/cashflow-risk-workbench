"""自造微型 xlsx 夹具生成器（源码交付，不只交手工截图）。

所有夹具都是明确的小号合成数据，与真实企业无关。
部分形态（公式缓存、空串、serial 60）openpyxl 写不出来，
用 patch_cell 直接改写包内 sheet XML 实现 —— 改写只发生在自己生成的夹具上。
"""
import datetime as dt
import os
import re
import shutil
import zipfile

import openpyxl

COMPANY_ROWS = [
    ("企业名称", "测试合成企业"),
    ("在职人数（人）", 320),
    ("上年度营业收入（万元）", 12800),
]

CASHFLOW_HEADER = ["月份", "月末日期", "流水类型", "经营性流入（万元）", "经营性流出（万元）"]
CASHFLOW_ROWS = [
    ["2026-01", dt.datetime(2026, 1, 31), "经营性", 180.87, 117.42],
    ["2026-02", dt.datetime(2026, 2, 28), "经营性", 184.35, 199.54],
    ["2026-03", dt.datetime(2026, 3, 31), "经营性", 174.57, 200.04],
]

AR_HEADER = ["单据编号", "应收金额（万元）", "已收金额（万元）", "到账状态"]


def _fill_company(ws, rows):
    ws["A1"] = "合成｜企业基本信息"
    ws["A2"], ws["B2"], ws["C2"] = "字段", "值", "说明"
    for i, (k, v) in enumerate(rows, start=3):
        ws[f"A{i}"] = k
        ws[f"B{i}"] = v


def _fill_sheet(ws, title, header, rows):
    ws["A1"] = title
    for j, h in enumerate(header, start=1):
        ws.cell(row=2, column=j, value=h)
    for i, row in enumerate(rows, start=3):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)


def build_workbook(path, *, company_rows=COMPANY_ROWS, cashflow_header=CASHFLOW_HEADER,
                   cashflow_rows=CASHFLOW_ROWS, include_cashflow=True,
                   ar_rows=None, date1904=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "企业基本信息"
    _fill_company(ws, company_rows)
    if include_cashflow:
        ws2 = wb.create_sheet("现金流水")
        _fill_sheet(ws2, "合成｜现金流水", cashflow_header, cashflow_rows)
    if ar_rows is not None:
        ws3 = wb.create_sheet("应收账款")
        _fill_sheet(ws3, "合成｜应收账款", AR_HEADER, ar_rows)
    if date1904:
        wb.epoch = dt.datetime(1904, 1, 1)
    wb.save(path)
    return path


def patch_cell(xlsx_path, sheet_index, addr, new_fragment):
    """把 xl/worksheets/sheet{sheet_index}.xml 中 addr 单元格整体替换为 new_fragment。

    new_fragment 必须是完整 <c .../> 或 <c ...>...</c> 元素。
    """
    sheetfile = f"xl/worksheets/sheet{sheet_index}.xml"
    tmp = xlsx_path + ".tmp"
    with zipfile.ZipFile(xlsx_path) as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    xml = data[sheetfile].decode("utf-8")
    pattern = re.compile(r'<c r="%s"[^>]*(?:/>|>.*?</c>)' % re.escape(addr), re.S)
    if not pattern.search(xml):
        raise RuntimeError(f"patch 目标不存在: {sheetfile}!{addr}")
    xml = pattern.sub(new_fragment.replace("\\", "\\\\"), xml, count=1)
    data[sheetfile] = xml.encode("utf-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, xlsx_path)


def set_formula(xlsx_path, sheet_index, addr, formula, cached):
    if cached is None:
        frag = f'<c r="{addr}"><f>{formula}</f></c>'
    else:
        frag = f'<c r="{addr}"><f>{formula}</f><v>{cached}</v></c>'
    patch_cell(xlsx_path, sheet_index, addr, frag)


def set_serial(xlsx_path, sheet_index, addr, serial):
    patch_cell(xlsx_path, sheet_index, addr, f'<c r="{addr}"><v>{serial}</v></c>')


def set_empty_string(xlsx_path, sheet_index, addr):
    patch_cell(xlsx_path, sheet_index, addr,
               f'<c r="{addr}" t="inlineStr"><is><t></t></is></c>')


def set_blank_cell(xlsx_path, sheet_index, addr):
    patch_cell(xlsx_path, sheet_index, addr, f'<c r="{addr}"/>')


def make_input_dir(root, name, **kw):
    d = os.path.join(root, name)
    os.makedirs(d, exist_ok=False)
    build_workbook(os.path.join(d, "fixture.xlsx"), **kw)
    return d
